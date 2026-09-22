"""Excel 批量导入引擎（数据进出方案 §2）。

解析结果由服务端暂存（`import_previews`），commit 只认 `preview_id`，
前端无法回传裸数据行绕过人工确认。所有入库仍走 `record()` 闸门。
"""

import csv
import io
import math
import re
import statistics
import unicodedata
import uuid
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from typing import Any

from charset_normalizer import from_bytes
from sqlalchemy import select
from sqlalchemy.orm import Session

from tin.compute.engine import compute_day
from tin.config import SHANGHAI
from tin.ingest.record import RecordError, record
from tin.models import ImportPreview, Indicator, Observation
from tin.schemas.caliber import Caliber
from tin.schemas.observation import ObservationIn

READY, REVISION, SKIP, ERROR = "ready", "revision", "skip", "error"
CODE_TAG = re.compile(r"\[([A-Za-z0-9_.]+)\]")
NARROW_HINTS = ("指标", "series", "代码")
VALUE_HINTS = ("数值", "value", "观测值", "值")
DATE_HINTS = ("日期", "date", "时间", "as_of", "时点")
FALLBACK_ENTRY_TIME = time(15, 0)
PREVIEW_TTL = timedelta(minutes=30)
# 预览耗时实测约 0.105 秒/千单元格，验收线是「返回 preview_id ≤ 15 秒」，
# 对应约 14 万单元格。取 12 万留出余量：12 万约 12.6 秒，payload 约 20MB。
# 再往上就不该走逐行预览了——没人逐行确认得了，该走数据商终端那条摘要确认通道（数据进出方案 §3）。
MAX_CELLS = 120_000

# 带这些口径维度的指标容易出现"只差一个字"的变体（40%/60% TC、实物吨/金属吨……）。
# 即便库里当前只登记了一个变体，按名称匹配也必须人工确认——表里那一列可能正是没登记的另一个。
VARIANT_PRONE = ("tc_grade", "weight_basis", "stock_scope", "spot_source", "import_code", "contract_kind")


class ImportError_(ValueError):
    pass


@dataclass
class ColumnSpec:
    col_key: str
    header: str
    series_id: str | None = None
    name: str | None = None
    unit: str | None = None
    match_kind: str = "none"  # code / name / none
    needs_choice: bool = False
    candidates: list[str] = field(default_factory=list)
    as_of_time: str | None = None  # "11:30"，仅当单元格只有日期时使用
    error: str | None = None


@dataclass
class RowItem:
    row_key: str
    col_key: str
    series_id: str
    as_of: str
    value: float
    category: str
    reason: str | None = None
    old_value: float | None = None
    time_filled: bool = False
    warn: str | None = None


def _norm(text: str) -> str:
    return re.sub(r"[\s（）()·\-_/]+", "", unicodedata.normalize("NFKC", str(text))).lower()


# ---------- 读表 ----------

def _decode_csv(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8")
    # 顺序按方案 §3.1.3：UTF-8 严格解码 → GB18030（中文环境主力编码）→ 探测器 → 兜底。
    # 探测器不能放在首位：短样本容易被猜成其它 CJK 编码，能解码但全是乱码且不报错。
    guess = from_bytes(data[:10240]).best()
    for encoding in ["utf-8", "gb18030"] + ([guess.encoding] if guess else []) + ["latin-1"]:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ImportError_("无法识别 CSV 文件编码，请另存为 UTF-8 或 GBK")


def read_table(data: bytes, filename: str) -> list[list[Any]]:
    """把上传文件读成二维数组，不做任何语义解释。"""
    name = filename.lower()
    if name.endswith(".csv"):
        rows = [list(r) for r in csv.reader(io.StringIO(_decode_csv(data)))]
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        rows = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
        wb.close()
    elif name.endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        rows = [[_xls_cell(sheet.cell(r, c), book.datemode) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]
    else:
        raise ImportError_("不支持的文件格式，仅允许上传 .xlsx、.xls、.csv")

    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ImportError_("文件里没有任何数据")
    cells = sum(len(r) for r in rows)
    if cells > MAX_CELLS:
        raise ImportError_(f"文件超出规模限制（最大支持 {MAX_CELLS:,} 单元格，本次 {cells:,}），请按时间分批导入")
    return rows


def _xls_cell(cell, datemode):
    import xlrd
    if cell.ctype == xlrd.XL_CELL_DATE:
        return datetime(*xlrd.xldate_as_tuple(cell.value, datemode))
    return cell.value


# ---------- 单元格解析 ----------

def parse_as_of(raw: Any, fill_time: time | None) -> tuple[datetime, bool]:
    """返回（带时区的时点, 是否由系统补齐了时刻）。只有日期时用指标的常规发布时刻补齐。"""
    if raw is None or str(raw).strip() == "":
        raise ValueError("时点为空")
    if isinstance(raw, datetime):
        naive, has_time = raw, (raw.hour or raw.minute)
    elif isinstance(raw, date):
        naive, has_time = datetime.combine(raw, time(0)), False
    elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
        # Excel 序列日期（1899-12-30 起算）
        naive = datetime(1899, 12, 30) + timedelta(days=float(raw))
        has_time = abs(float(raw) - int(raw)) > 1e-9
    else:
        text = str(raw).strip().replace("/", "-")
        naive, has_time = None, False
        for fmt, with_time in (("%Y-%m-%d %H:%M:%S", True), ("%Y-%m-%d %H:%M", True),
                               ("%Y-%m-%dT%H:%M", True), ("%Y-%m-%d", False), ("%Y%m%d", False)):
            try:
                naive, has_time = datetime.strptime(text, fmt), with_time
                break
            except ValueError:
                continue
        if naive is None:
            raise ValueError(f"无法识别时间文本「{raw}」")
    if has_time:
        return naive.replace(tzinfo=SHANGHAI), False
    filled = datetime.combine(naive.date(), fill_time or FALLBACK_ENTRY_TIME)
    return filled.replace(tzinfo=SHANGHAI), True


def parse_value(raw: Any, unit: str | None = None) -> float:
    if raw is None or str(raw).strip() == "":
        raise ValueError("空值")
    if isinstance(raw, bool):
        raise ValueError(f"非数值内容「{raw}」")
    if isinstance(raw, (int, float)):
        value = float(raw)
    else:
        text = str(raw).strip().replace(",", "").replace("，", "")
        percent = text.endswith("%")
        text = re.sub(r"[¥$￥元吨手点%\s]|万吨|金属吨|实物吨", "", text)
        if not re.fullmatch(r"-?\d+(\.\d+)?", text):
            raise ValueError(f"非数值内容「{raw}」")
        value = float(text)
        if percent and unit != "%":
            value /= 100
    if not math.isfinite(value):
        raise ValueError(f"非有限数值「{raw}」")
    return value


# ---------- 表头匹配 ----------

def match_header(header: str, indicators: list[Indicator]) -> ColumnSpec:
    """优先级 1：表头含 [series_id] 精确对齐；优先级 2：名称模糊匹配，受品位防呆约束。"""
    spec = ColumnSpec(col_key="", header=str(header))
    by_id = {i.series_id: i for i in indicators}

    tag = CODE_TAG.search(str(header))
    if tag:
        ind = by_id.get(tag.group(1))
        if ind is None:
            spec.error = f"表头代码 {tag.group(1)} 未在指标登记表中找到"
            return spec
        spec.series_id, spec.match_kind = ind.series_id, "code"
        return _finish(spec, ind)

    target = _norm(header)
    hits = [i for i in indicators
            if _norm(i.name) == target or (i.vendor_code and _norm(i.vendor_code) == target)]
    if not hits:
        hits = [i for i in indicators
                if target and (target in _norm(i.name) or _norm(i.name) in target)]
    if not hits:
        spec.error = "系统未登记该指标"
        return spec
    if len(hits) > 1:
        spec.needs_choice = True
        spec.candidates = [i.series_id for i in hits]
        spec.error = "名称匹配到多个指标，请手动选定"
        return spec

    ind = hits[0]
    spec.series_id, spec.match_kind = ind.series_id, "name"
    spec = _finish(spec, ind)
    if spec.error is None and any(ind.caliber.get(d) for d in VARIANT_PRONE):
        # 库里可能只登记了 40% 这一个变体，表里却是 60%：名称匹配不足以确认口径
        spec.needs_choice = True
        spec.candidates = [i.series_id for i in indicators if i.category == ind.category]
        spec.error = "该指标存在易混口径（品位/重量口径等），请手动确认后再导入"
    return spec


def _finish(spec: ColumnSpec, ind: Indicator) -> ColumnSpec:
    spec.name, spec.unit = ind.name, ind.unit
    spec.as_of_time = (ind.default_entry_time or FALLBACK_ENTRY_TIME).strftime("%H:%M")
    if ind.status != "可用":
        spec.error = "指标已停用"
    elif ind.fetch_mode != "manual":
        spec.error = "自动采集指标不接受人工导入，防止篡改交易所或官方权威行情"
    return spec


# ---------- 数量级防呆 ----------

def magnitude_warning(session: Session, series_id: str, value: float) -> str | None:
    history = [abs(v) for v in session.scalars(
        select(Observation.value).where(Observation.series_id == series_id)
        .order_by(Observation.as_of.desc()).limit(20)) if v]
    if len(history) < 5:
        return None
    median = statistics.median(history)
    if median == 0 or value == 0:
        return None
    if abs(math.log10(abs(value)) - math.log10(median)) >= 2.0:
        return (f"数值 {value:,.4g} 与历史中位数 {median:,.4g} 相差达 2 个数量级，"
                f"请核对是否发生单位混淆（如元与万元、吨与万吨）")
    return None


# ---------- 形态识别与预览 ----------

def detect_shape(header: list[Any]) -> str:
    cells = [_norm(c) for c in header if c not in (None, "")]
    has_series = any(any(h in c for h in NARROW_HINTS) for c in cells)
    has_value = any(any(h in c for h in VALUE_HINTS) for c in cells)
    return "narrow" if has_series and has_value else "wide"


def _find(header: list[Any], hints: tuple[str, ...]) -> int | None:
    for i, cell in enumerate(header):
        if cell and any(h in _norm(cell) for h in hints):
            return i
    return None


def _classify(session: Session, series_id: str, as_of: datetime, value: float) -> tuple[str, float | None]:
    latest = session.scalars(
        select(Observation).where(Observation.series_id == series_id, Observation.as_of == as_of)
        .order_by(Observation.revision.desc()).limit(1)).first()
    if latest is None:
        return READY, None
    if math.isclose(latest.value, value, rel_tol=1e-12):
        return SKIP, latest.value
    return REVISION, latest.value


def build_preview(session: Session, data: bytes, filename: str, actor: str, variety: str) -> dict:
    rows = read_table(data, filename)
    header = rows[0]
    indicators = list(session.scalars(select(Indicator)))
    columns: list[ColumnSpec] = []
    items: list[RowItem] = []

    if detect_shape(header) == "narrow":
        i_series = _find(header, NARROW_HINTS)
        i_as_of = _find(header, DATE_HINTS)
        i_value = _find(header, VALUE_HINTS)
        i_note = _find(header, ("来源", "备注", "凭据"))
        if i_series is None or i_as_of is None or i_value is None:
            raise ImportError_("窄表缺少必要列：指标代码 / 数据时点 / 观测数值")
        by_key = {}
        for r, row in enumerate(rows[1:], start=2):
            raw_series = row[i_series] if i_series < len(row) else None
            spec = by_key.get(str(raw_series))
            if spec is None:
                spec = match_header(str(raw_series or ""), indicators)
                spec.col_key = f"c{len(by_key)}"
                by_key[str(raw_series)] = spec
                columns.append(spec)
            _append(session, items, spec, r, row[i_as_of] if i_as_of < len(row) else None,
                    row[i_value] if i_value < len(row) else None,
                    row[i_note] if i_note is not None and i_note < len(row) else None)
    else:
        i_as_of = _find(header, DATE_HINTS) or 0
        for c, cell in enumerate(header):
            if c == i_as_of or cell in (None, ""):
                continue
            spec = match_header(str(cell), indicators)
            spec.col_key = f"c{c}"
            columns.append(spec)
        for r, row in enumerate(rows[1:], start=2):
            when = row[i_as_of] if i_as_of < len(row) else None
            for spec in columns:
                c = int(spec.col_key[1:])
                _append(session, items, spec, r, when, row[c] if c < len(row) else None, None)

    payload = {"variety": variety, "columns": [asdict(c) for c in columns],
               "rows": [asdict(i) for i in items]}
    now = datetime.now(timezone.utc)
    preview = ImportPreview(id=str(uuid.uuid4()), variety=variety, actor=actor, filename=filename,
                            payload=payload, created_at=now, expires_at=now + PREVIEW_TTL)
    session.execute(ImportPreview.__table__.delete().where(ImportPreview.expires_at < now))
    session.add(preview)
    session.commit()
    return {"preview_id": preview.id, **summarize(payload)}


def _append(session: Session, items: list[RowItem], spec: ColumnSpec, row_no: int,
            raw_when: Any, raw_value: Any, raw_note: Any) -> None:
    if raw_value is None or str(raw_value).strip() == "":
        return
    key = f"r{row_no}-{spec.col_key}"
    if spec.error and not spec.needs_choice:
        items.append(RowItem(key, spec.col_key, spec.series_id or "", "", 0.0, ERROR,
                             f"第 {row_no} 行 [{spec.header}]：{spec.error}"))
        return
    fill = time.fromisoformat(spec.as_of_time) if spec.as_of_time else None
    try:
        as_of, filled = parse_as_of(raw_when, fill)
    except ValueError as e:
        items.append(RowItem(key, spec.col_key, spec.series_id or "", "", 0.0, ERROR,
                             f"第 {row_no} 行 [时点]：{e}"))
        return
    try:
        value = parse_value(raw_value, spec.unit)
    except ValueError as e:
        items.append(RowItem(key, spec.col_key, spec.series_id or "", as_of.isoformat(), 0.0, ERROR,
                             f"第 {row_no} 行 [{spec.header}]：{e}"))
        return

    if spec.needs_choice or not spec.series_id:
        items.append(RowItem(key, spec.col_key, spec.series_id or "", as_of.isoformat(), value,
                             ERROR, f"第 {row_no} 行 [{spec.header}]：{spec.error}", time_filled=filled))
        return
    category, old = _classify(session, spec.series_id, as_of, value)
    items.append(RowItem(key, spec.col_key, spec.series_id, as_of.isoformat(), value, category,
                         old_value=old, time_filled=filled,
                         warn=magnitude_warning(session, spec.series_id, value)))


def summarize(payload: dict) -> dict:
    counts = {c: 0 for c in (READY, REVISION, SKIP, ERROR)}
    for row in payload["rows"]:
        counts[row["category"]] += 1
    return {"ready_count": counts[READY], "revision_count": counts[REVISION],
            "skipped_count": counts[SKIP], "error_count": counts[ERROR],
            "needs_choice": [c["col_key"] for c in payload["columns"] if c["needs_choice"]],
            "columns": payload["columns"], "rows": payload["rows"]}


# ---------- 入库 ----------

def commit_preview(session: Session, preview_id: str, selected_row_keys: list[str],
                   column_overrides: dict[str, dict], actor: str, default_note: str) -> dict:
    preview = session.get(ImportPreview, preview_id)
    if preview is None:
        raise ImportError_("预览不存在或已被清理，请重新上传文件")
    if preview.expires_at < datetime.now(timezone.utc):
        raise ImportError_("预览已过期（超过 30 分钟），请重新上传文件")

    columns = {c["col_key"]: dict(c) for c in preview.payload["columns"]}
    indicators = {i.series_id: i for i in session.scalars(select(Indicator))}

    # 列改绑必须重新校验：新指标要存在、可用、且是人工录入类
    for col_key, override in (column_overrides or {}).items():
        col = columns.get(col_key)
        if col is None:
            raise ImportError_(f"未知的列 {col_key}")
        if sid := override.get("series_id"):
            ind = indicators.get(sid)
            if ind is None:
                raise ImportError_(f"改绑的指标未登记：{sid}")
            if ind.status != "可用":
                raise ImportError_(f"改绑的指标已停用：{sid}")
            if ind.fetch_mode != "manual":
                raise ImportError_(f"{sid} 为自动采集指标，不接受人工导入")
            col.update({"series_id": sid, "name": ind.name, "unit": ind.unit,
                        "needs_choice": False, "error": None})
        if at := override.get("as_of_time"):
            col["as_of_time"] = at
        if note := override.get("note"):
            col["note"] = note

    selected = set(selected_row_keys)
    committed, revisions, stale, dates, warnings = 0, 0, [], set(), []
    for raw in preview.payload["rows"]:
        if raw["row_key"] not in selected or raw["category"] == SKIP:
            continue
        col = columns[raw["col_key"]]
        series_id = col.get("series_id")
        if not series_id or col.get("error"):
            stale.append({"row_key": raw["row_key"], "reason": col.get("error") or "该列未完成口径确认"})
            continue

        as_of = datetime.fromisoformat(raw["as_of"])
        if col.get("as_of_time") and raw["time_filled"]:
            hh, mm = (int(x) for x in col["as_of_time"].split(":"))
            as_of = as_of.replace(hour=hh, minute=mm)
        value = raw["value"]

        # 并发陈旧检测：预览与确认之间库内可能已被写入
        category, old = _classify(session, series_id, as_of, value)
        if category == SKIP:
            continue
        # 改绑、或预览时该列尚未完成口径确认（行被归为异常），都要按当前状态重新判定
        rebound = col.get("series_id") != raw["series_id"] or raw["category"] == ERROR
        if rebound:
            # 若重判后变成修订，用户勾选时并不知情，仍要拦下
            if category == REVISION:
                stale.append({"row_key": raw["row_key"], "series_id": series_id, "old_value": old,
                              "reason": "改绑后该时点已有不同数值，请重新预览确认"})
                continue
            if warn := magnitude_warning(session, series_id, value):
                warnings.append({"row_key": raw["row_key"], "warn": warn})
        elif category != raw["category"]:
            stale.append({"row_key": raw["row_key"], "series_id": series_id,
                          "as_of": raw["as_of"], "old_value": old,
                          "reason": "预览后库内数据已变化，已拦截改写"})
            continue

        ind = indicators[series_id]
        note = col.get("note") or default_note
        if raw["time_filled"]:
            note = f"[截点补齐: {col.get('as_of_time') or ''}] {note}".strip()
        try:
            row = record(session, ObservationIn(
                series_id=series_id, value=value, as_of=as_of,
                caliber=Caliber.model_validate(ind.caliber), source="人工",
                entered_by=actor, note=note))
        except RecordError as e:
            stale.append({"row_key": raw["row_key"], "reason": str(e)})
            continue
        if row is not None:
            committed += 1
            revisions += 1 if row.revision else 0
            dates.add(as_of.astimezone(SHANGHAI).date())
    session.commit()

    # 派生公式只读 observations，重算顺序与结果无关；按升序仅为日志可读
    for d in sorted(dates):
        compute_day(session, d, preview.variety)
    return {"committed": committed, "revisions": revisions, "stale_rows": stale,
            "warnings": warnings, "recalculated_dates": [d.isoformat() for d in sorted(dates)]}
