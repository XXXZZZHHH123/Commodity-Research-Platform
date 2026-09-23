"""数据商终端导出工作簿的整体导入（SMM、钢联 Mysteel）。

两家的导出长得不一样，但结构同类：**若干行元数据在上、日期索引的数据在下**，
而且都给了**稳定的供应商编码**。所以这里只写一份读取器，各家的差异收在 `Vendor` 声明里。

    SMM                         钢联 Mysteel
    R1 指标名称                 R1 钢联数据            ← 横幅
    R2 指标Id                   R2 指标名称
    R3 单位                     R3 单位
    R4 频率                     R4 指标编码
    R5+ 数据                    R5 频度
                                R6 指标描述
                                R7 更新时间（部分表没有）
                                R8+ 数据

表头行数**不固定**（钢联 Sheet1 就少一行），所以数据起点靠「首列是日期」判断，不靠行号。

供应商编码是这条通道的地基：它稳定、唯一、与显示名无关，因此不需要任何名称猜测。
`indicators.vendor_code` 本来就是为此准备的字段。
"""

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from tin.caliber.dictionary import (
    ContractKind,
    PriceType,
    SpotSource,
    StockScope,
    TCGrade,
    TimeType,
    WeightBasis,
)
from sqlalchemy import select

from tin.config import SHANGHAI
from tin.schemas.caliber import Caliber
from tin.schemas.observation import IndicatorSpec, ObservationIn


@dataclass
class Series:
    vendor: str
    code: str
    name: str
    unit: str
    frequency: str
    sheet: str
    description: str | None = None
    points: dict[date, float] = field(default_factory=dict)


class VendorFormatError(ValueError):
    pass


# 自动登记但无人复核过口径的指标，`owner` 标成这个值
UNVERIFIED = "自动登记·未核验"


@dataclass(frozen=True)
class Vendor:
    key: str
    label: str                       # 显示名，写进 observations.source
    source_url: str
    prefix: str                      # series_id 前缀
    labels: dict[str, str]           # 行首标签 → 字段名
    required: tuple[str, ...]        # 探测用：这些标签都出现才算命中
    missing: tuple[str, ...]         # 缺失值标记
    overrides: dict[str, str]        # 供应商编码 → 已登记 series_id
    entry_time: dict[str, time]      # 原始发布方 → 常规发布时刻
    publisher: Callable[[str], str]
    caliber: Callable[["Series"], Caliber]
    discontinued: str | None = None  # 名称前缀，命中则不导入

    def series_id(self, code: str) -> str:
        return self.overrides.get(code, f"{self.prefix}{code}")


# ---------- 通用读取 ----------

_SCAN_ROWS = 12  # 元数据行最多扫这么多行，再往下必然是数据


def _header_of(ws, vendors: tuple[Vendor, ...]) -> tuple[Vendor, dict[str, list], int] | None:
    """扫出这张表的元数据行与数据起点。认不出格式就返回 None，整张跳过，不猜。"""
    known = {lab for v in vendors for lab in v.labels}
    head: dict[str, list] = {}
    for offset, row in enumerate(ws.iter_rows(max_row=_SCAN_ROWS, values_only=True)):
        first = row[0]
        if isinstance(first, datetime):
            for vendor in vendors:
                if all(lab in head for lab in vendor.required):
                    return vendor, head, offset
            return None
        label = str(first).strip() if first is not None else ""
        if label in known:
            head[label] = list(row)
    return None


def read_workbook(data_or_path) -> list[Series]:
    """读出全部序列。表头认不出的工作表整张跳过。"""
    from openpyxl import load_workbook

    src = io.BytesIO(data_or_path) if isinstance(data_or_path, bytes) else data_or_path
    wb = load_workbook(src, read_only=True, data_only=True)
    out: list[Series] = []
    seen: set[str] = set()
    try:
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            found = _header_of(ws, VENDORS)
            if found is None:
                continue
            vendor, head, start = found
            get = lambda lab, i: (  # noqa: E731
                head[lab][i] if lab in head and i < len(head[lab]) else None)
            fields = {v: k for k, v in vendor.labels.items()}
            cols: dict[int, Series] = {}
            names = head[fields["name"]]
            for i in range(1, len(names)):
                code = get(fields["code"], i)
                code = str(code).strip() if code else ""
                if not names[i] or not code or code in seen:
                    continue  # 同一编码重复出现（短名 + 全名）时只取先出现的那份
                seen.add(code)
                desc = get(fields.get("description", ""), i)
                cols[i] = Series(
                    vendor=vendor.key, code=code, name=str(names[i]).strip(), sheet=sheet,
                    unit=str(get(fields["unit"], i) or "").strip(),
                    frequency=str(get(fields["frequency"], i) or "").strip(),
                    description=str(desc).strip() if desc and str(desc).strip() != "·" else None,
                )
            for row in ws.iter_rows(min_row=start + 1, values_only=True):
                day = row[0]
                if not isinstance(day, datetime):
                    continue
                day = day.date()
                for i, series in cols.items():
                    if i >= len(row):
                        continue
                    value = row[i]
                    if value is None or (isinstance(value, str)
                                         and value.strip() in vendor.missing):
                        continue
                    try:
                        series.points[day] = float(value)
                    except (TypeError, ValueError):
                        continue  # 文本备注，不是数值序列
            out.extend(cols.values())
    finally:
        wb.close()
    if not out:
        raise VendorFormatError(f"{getattr(data_or_path, 'name', '文件')} 里没有可识别的数据商终端格式工作表")
    if not any(s.points for s in out):
        _explain_empty(src, len(out))
    return out


def _explain_empty(src, series_count: int) -> None:
    """一条数据都没读出来时，先弄清是不是公式没算过，再报错。

    Excel 插件（数据终端的自动更新加载项）写进单元格的是公式，公式的**计算结果**只有
    被 Excel 打开并保存过才会缓存进文件。openpyxl 不算公式，读到的就是 None——
    于是整本文件会被当成"全是缺失值"静默导入 0 条。这种失败必须吵，不能安静。
    """
    from openpyxl import load_workbook

    if hasattr(src, "seek"):
        src.seek(0)
    wb = load_workbook(src, read_only=True, data_only=False)
    try:
        formulas = sum(
            1 for sheet in wb.sheetnames
            for row in wb[sheet].iter_rows(max_row=200, values_only=True)
            for c in row
            if isinstance(c, str) and c.startswith("=") and not c.startswith("=NA(")
        )
    finally:
        wb.close()
    if formulas:
        raise VendorFormatError(
            f"识别出 {series_count} 条序列但一个数据点都没有，文件里有 {formulas} 处公式尚未计算。"
            "数据终端的 Excel 插件写入的是公式，计算结果要用 Excel 打开并保存后才会存进文件。"
            "请在 Excel 里刷新并保存一次再上传。")
    raise VendorFormatError(f"识别出 {series_count} 条序列，但一个数据点都没有——请确认文件内容非空")


def detect(data: bytes) -> Vendor | None:
    """只看前若干行的首列标签判断供应商。要在单元格上限之前做，所以不整本读。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - 打不开就不是，交回原有通道去报错
        return None
    try:
        for sheet in wb.sheetnames:
            labels = set()
            for row in wb[sheet].iter_rows(max_row=_SCAN_ROWS, max_col=1, values_only=True):
                if row and row[0] is not None:
                    labels.add(str(row[0]).strip())
            for vendor in VENDORS:
                if all(lab in labels for lab in vendor.required):
                    return vendor
        return None
    finally:
        wb.close()


def resolve_ids(session, series_list: list[Series]) -> tuple[dict[str, str], list[str]]:
    """编码 → series_id。**库里的 `vendor_code` 绑定优先于代码里的 `OVERRIDES`。**

    只查 `OVERRIDES` 的话，研究员在库里把某个编码登记到一个可读的 series_id 上之后，
    导入仍会另建一个 `SMM.<编码>`，同一条序列无声裂成两个——而且没有任何提示。
    人在库里做的绑定是权威的，代码里的 `OVERRIDES` 只是尚未登记时的引导。

    返回 (映射, 冲突说明)。同一编码被登记到多个指标时不替人选，报出来。
    """
    from tin.models import Indicator

    codes = {s.code for s in series_list}
    rows = session.scalars(select(Indicator).where(Indicator.vendor_code.in_(codes))).all()
    by_code: dict[str, list[str]] = {}
    for row in rows:
        by_code.setdefault(row.vendor_code, []).append(row.series_id)

    mapping, conflicts = {}, []
    for series in series_list:
        vendor = vendor_of(series)
        bound = by_code.get(series.code, [])
        if len(bound) > 1:
            conflicts.append(f"编码 {series.code} 同时登记在 {'、'.join(sorted(bound))}，请先合并")
            mapping[series.code] = sorted(bound)[0]
        elif bound:
            mapping[series.code] = bound[0]
        else:
            mapping[series.code] = vendor.series_id(series.code)
    return mapping, conflicts


def column_cache(session, series_list: list[Series], ids: dict[str, str],
                 *, count_hit: bool = True) -> dict:
    """编码通道接不住的列，查一次列映射学习缓存（`ingest/column_map.py`）。

    「接不住」= 这条编码在库里没有 `vendor_code` 绑定，退回去的 `SMM.<编码>` 也还没
    登记过——正是过去每次都要把文件拿给 AI 认一遍的那些列。缓存命中就就地改写
    `ids`（于是后面照常走已登记指标那条路，口径以登记为准）；没命中就原样上报，
    **不猜**：这里是兜底通道，兜不住时该让人来看，而不是自作主张。

    指纹按**工作表**算：一个 sheet 就是一张报表，换日期导出时表头不变、指纹不变；
    一本工作簿里多导了或少导了一张表，也不会牵连其它表的映射。

    返回 `{"fingerprints": {sheet: 指纹}, "mapped": {编码: series_id}, "unmapped": [...]}`，
    `unmapped` 的每一项都带齐 sheet / 指纹 / 列名 / 编码，调用方拿着就能直接调
    `column_map.confirm()`，不必再回头重算一次指纹。
    """
    from tin.ingest import column_map
    from tin.models import Indicator

    by_sheet: dict[str, list[Series]] = {}
    for series in series_list:
        by_sheet.setdefault(series.sheet, []).append(series)
    # 一次问清哪些 series_id 已登记，别在几千条序列上逐条 get
    known = set(session.scalars(select(Indicator.series_id).where(
        Indicator.series_id.in_({ids[s.code] for s in series_list}))))

    result: dict = {"fingerprints": {}, "mapped": {}, "unmapped": []}
    for sheet, group in by_sheet.items():
        fp = column_map.fingerprint([s.name for s in group])
        result["fingerprints"][sheet] = fp
        pending = [s for s in group if ids[s.code] not in known]
        if not pending:
            continue
        mapped, missing = column_map.lookup(
            session, fp, [s.name for s in pending], count_hit=count_hit)
        for series in pending:
            hit = mapped.get(series.name)
            if hit:
                ids[series.code] = hit
                result["mapped"][series.code] = hit
        unresolved = set(missing)
        result["unmapped"] += [
            {"sheet": sheet, "fingerprint": fp, "column_name": s.name, "vendor_code": s.code}
            for s in pending if s.name in unresolved
        ]
    return result


def vendor_of(series: Series) -> Vendor:
    return next(v for v in VENDORS if v.key == series.vendor)


def discontinued(series: Series) -> bool:
    mark = vendor_of(series).discontinued
    return bool(mark) and series.name.startswith(mark)


# ---------- SMM ----------

_SMM_PRICE = (("结算价", PriceType.结算价), ("收盘价", PriceType.收盘价),
              ("平均价", PriceType.均价), ("均价", PriceType.均价),
              ("开盘价", PriceType.开盘价), ("最高价", PriceType.最高价),
              ("最低价", PriceType.最低价))
_SMM_STOCK = (("仓单", StockScope.交易所注册仓单), ("社会库存", StockScope.社会库存),
              ("库存", StockScope.交易所库存))
_SMM_CONTRACT = (("主力合约", ContractKind.主力合约), ("连二合约", ContractKind.连二连续))
_SMM_TIME = {"SHFE": time(15, 0), "LME": time(19, 0), "SMM": time(11, 30)}


def _smm_publisher(name: str) -> str:
    head = name.removeprefix("（停）").split(":")[0].strip()
    return head if head in _SMM_TIME else "SMM"


def _smm_spot_source(name: str) -> SpotSource | None:
    """只认现货价序列。个旧不在口径字典里，留空并标注，不塞字典外的值。"""
    if "1#锡" not in name:
        return None
    if "长江" in name:
        return SpotSource.长江有色
    return None if "个旧" in name else SpotSource.SMM


def _smm_caliber(series: Series) -> Caliber:
    name, who = series.name, _smm_publisher(series.name)
    kwargs: dict = {"time_type": TimeType.交易时点 if who in ("SHFE", "LME") else TimeType.发布时点}
    for token, value in _SMM_PRICE:
        if token in name:
            kwargs["price_type"] = value
            break
    for token, value in _SMM_CONTRACT:
        if token in name:
            kwargs["contract_kind"] = value
            break
    for token, value in _SMM_STOCK:
        if token in name:
            kwargs["stock_scope"] = value
            break
    source = _smm_spot_source(name)
    if source is not None:
        kwargs["spot_source"] = source
    if series.unit in (WeightBasis.金属吨, WeightBasis.实物吨):
        kwargs["weight_basis"] = series.unit

    note = f"SMM 终端 {series.code}；原始发布方 {who}"
    if "分仓库" in name or "分地区" in name:
        note += "；分项口径，不可与总计相加混用"
    if "个旧" in name:
        note += "；现货来源为个旧市场，口径字典暂无该枚举，spot_source 留空"
    return Caliber(note=note, **kwargs)


SMM = Vendor(
    key="SMM", label="SMM 终端",
    source_url="https://www.smm.cn/（终端导出，订阅授权数据，禁止外发）",
    prefix="SMM.",
    labels={"指标名称": "name", "指标Id": "code", "单位": "unit", "频率": "frequency"},
    required=("指标名称", "指标Id", "单位", "频率"),
    missing=("", "-"),
    overrides={"s20002960": "SMM.SN.spot.1"},
    entry_time=_SMM_TIME, publisher=_smm_publisher, caliber=_smm_caliber,
    discontinued="（停）",
)


# ---------- 钢联 Mysteel ----------

_MS_TIME = {"SHFE": time(15, 0), "LME": time(19, 0), "Mysteel": time(15, 0)}
_MS_PRICE = (("结算价", PriceType.结算价), ("汇总均价", PriceType.均价), ("汇总价格", PriceType.均价),
             ("最高成交价", PriceType.最高价), ("最低成交价", PriceType.最低价))
_MS_CONTRACT = (("主力合约", ContractKind.主力合约), ("连续合约", ContractKind.主力连续))
_GRADE = re.compile(r"(\d{2})%Sn")


def _ms_publisher(name: str) -> str:
    head = name.split("：")[0].strip()
    return head if head in _MS_TIME else "Mysteel"


def _ms_caliber(series: Series) -> Caliber:
    name, who = series.name, _ms_publisher(series.name)
    kwargs: dict = {"time_type": TimeType.交易时点 if who in ("SHFE", "LME") else TimeType.发布时点}

    for token, value in _MS_PRICE:
        if token in name:
            kwargs["price_type"] = value
            break
    for token, value in _MS_CONTRACT:
        if token in name:
            kwargs["contract_kind"] = value
            break

    grade = _GRADE.search(name)
    if grade:
        try:
            kwargs["tc_grade"] = TCGrade(f"{int(grade.group(1))}度")
        except ValueError:
            pass  # 字典里没有这个品位，留空好过塞个字典外的值

    if "库存" in name:
        # 钢联的锡锭库存是它自己的社会库存调研——ID01517441 已按社会库存登记可作佐证；
        # LME 的期货库存是交易所库存，两者不可混算。
        kwargs["stock_scope"] = StockScope.交易所库存 if who == "LME" else StockScope.社会库存

    note = f"钢联 {series.code}；原始发布方 {who}"
    if series.description:
        note += f"；{series.description[:60]}"
    if "加工费" in name:
        note += "；加工费不是价格类型，price_type 留空"
    return Caliber(note=note, **kwargs)


MYSTEEL = Vendor(
    key="MYSTEEL", label="钢联终端",
    source_url="https://www.mysteel.com/（终端导出，订阅授权数据，禁止外发）",
    prefix="MYSTEEL.",
    labels={"指标名称": "name", "指标编码": "code", "单位": "unit",
            "频度": "frequency", "指标描述": "description"},
    required=("指标名称", "指标编码", "单位", "频度"),
    missing=("", "-", "#N/A", "NaN"),
    overrides={
        "ID01538256": "MYSTEEL.SN.TC.YN40",        # 云南 40% 加工费，判断的多头证据 bull1
        "ID01517441": "MYSTEEL.SN.stock.social",   # 锡锭库存，证伪条件 fl_social_stock_2w
        "FU00082529": "ICDX.SN.volume",            # 印尼锡锭成交量，白名单指标
    },
    entry_time=_MS_TIME, publisher=_ms_publisher, caliber=_ms_caliber,
)


VENDORS: tuple[Vendor, ...] = (SMM, MYSTEEL)


# ---------- 指标登记与观测 ----------

_CATEGORY_RULES = (
    ("仓单", "库存与流通"), ("库存", "库存与流通"),
    ("矿端平衡", "矿端"), ("锡精矿", "矿端"), ("锡矿", "矿端"),
    ("锡市平衡", "供需平衡"), ("供需平衡", "供需平衡"), ("平衡", "供需平衡"),
    ("强度指数", "情绪"), ("强弱指数", "情绪"), ("价格指数", "价格"),
    ("消费结构", "需求"), ("消费", "需求"), ("消耗量", "需求"),
    ("开工率", "需求"), ("产量", "矿端"),
    ("溢价", "价格"), ("升贴水", "价格"), ("价差", "价格"), ("加工费", "矿端"),
    ("价", "价格"), ("成交量", "价格"), ("持仓量", "价格"),
)


def derive_category(name: str) -> str:
    for token, category in _CATEGORY_RULES:
        if token in name:
            return category
    return "其他"


def derive_caliber(series: Series) -> Caliber:
    """按各家的命名规则推导口径。"""
    return vendor_of(series).caliber(series)


def caliber_for(series: Series, registered=None) -> Caliber:
    """已登记的指标以登记口径为准，未登记的才从名称推导。

    登记是人按供应商编码做的绑定，比我按名称的推导权威；单条人工录入
    （`jobs enter`）读的也是 `ind.caliber`，这里保持一致。
    """
    if registered is not None:
        return Caliber.model_validate(registered.caliber)
    return vendor_of(series).caliber(series)


def entry_time_for(series: Series, registered=None) -> time:
    if registered is not None and registered.default_entry_time:
        return registered.default_entry_time
    vendor = vendor_of(series)
    return vendor.entry_time.get(vendor.publisher(series.name), time(15, 0))


def derive_spec(series: Series, variety: str = "SN", series_id: str | None = None) -> IndicatorSpec:
    vendor = vendor_of(series)
    return IndicatorSpec(
        series_id=series_id or vendor.series_id(series.code),
        name=series.name[:120], variety=variety,
        category=derive_category(series.name),
        caliber=vendor.caliber(series),
        unit=(series.unit or "-")[:16],
        source=vendor.label, source_url=vendor.source_url,
        frequency=(series.frequency or "日")[:8],
        fetch_mode="manual",
    )


def observations(series: Series, entered_by: str, registered=None,
                 variety: str = "SN", series_id: str | None = None) -> list[ObservationIn]:
    """把一条序列的所有点转成待入库观测。

    日期单元格只有日期没有时刻，按发布方的常规发布时刻补齐，并在 note 里写明补齐规则——
    系统不替人编造时点，但补齐规则必须显式可查。
    """
    vendor = vendor_of(series)
    caliber = caliber_for(series, registered)
    at = entry_time_for(series, registered)
    note = (f"{vendor.label}导出 {series.code}；"
            f"时刻按 {vendor.publisher(series.name)} 常规发布时刻 {at:%H:%M} 补齐")
    sid = series_id or vendor.series_id(series.code)
    return [
        ObservationIn(series_id=sid, value=value,
                      as_of=datetime.combine(day, at, SHANGHAI), caliber=caliber,
                      source=vendor.label, source_url=vendor.source_url,
                      entered_by=entered_by, note=note)
        for day, value in sorted(series.points.items())
    ]


# ---------- 摘要与入库 ----------

def summarize(session, data_or_path, variety: str = "SN") -> dict:
    """确认页要的批次级信息。

    几十万条时人能判断的是「多少新指标、跨度多久、有没有口径冲突」，
    不是「第 137,204 行的值对不对」。
    """
    from collections import Counter

    from tin.models import Indicator

    everything = read_workbook(data_or_path)
    active = [s for s in everything if not discontinued(s)]

    ids, id_conflicts = resolve_ids(session, active)
    # 确认页只是看看，不计命中数——`hit_count` 记的是真正入库的次数
    cache = column_cache(session, active, ids, count_hit=False)
    new, reused, notes, mismatched = [], [], [], []
    for series in active:
        spec = derive_spec(series, variety, series_id=ids[series.code])
        existing = session.get(Indicator, spec.series_id)
        if existing is None:
            new.append(spec)
            continue
        reused.append({"series_id": spec.series_id, "name": existing.name})
        if existing.frequency and series.frequency and existing.frequency != series.frequency:
            mismatched.append(f"{spec.series_id}：登记为{existing.frequency}频，文件为{series.frequency}频")
        # 已登记的以登记口径为准（人按供应商编码做的绑定比按名称的推导权威），
        # 但两者不一致时要说出来：要么我的推导规则有问题，要么这条登记绑错了编码。
        derived = derive_caliber(series).dump()
        booked = Caliber.model_validate(existing.caliber).dump()
        differing = [k for k in set(derived) | set(booked)
                     if k != "note" and derived.get(k) != booked.get(k)]
        if differing:
            notes.append(f"{spec.series_id}：" + "；".join(
                f"{k} 登记为 {booked.get(k) or '空'}，按名称推导为 {derived.get(k) or '空'}"
                for k in sorted(differing)))

    days = [d for s in active for d in (min(s.points, default=None), max(s.points, default=None)) if d]
    vendors = sorted({vendor_of(s).label for s in everything})
    # 年频/季频按期末标注，于是会出现晚于今天的点：2026 年的年度均值标在 2026-12-31。
    # 这不是错，但那一年还没过完，值是阶段性的——导入前要让人知道有多少这种点。
    today = date.today()
    future = sum(1 for s in active for d in s.points if d > today)
    return {
        "mode": "vendor_terminal",
        "vendor": "、".join(vendors),
        "series_total": len(everything),
        "new_indicators": len(new),
        "reused_indicators": reused,
        "skipped_discontinued": len(everything) - len(active),
        "points": sum(len(s.points) for s in active),
        "first_date": min(days).isoformat() if days else None,
        "last_date": max(days).isoformat() if days else None,
        "categories": dict(Counter(spec.category for spec in new).most_common()),
        # 以登记为准，这里只是把差异摆出来给人看，不阻断
        "future_points": future,
        "code_conflicts": id_conflicts,
        "caliber_notes": notes,
        "frequency_mismatches": mismatched,
        # 列映射学习缓存：命中的自动走，没命中的列在这儿等人确认一次，下次就不用看了
        "column_fingerprints": cache["fingerprints"],
        "column_cache_hits": len(cache["mapped"]),
        "unmapped_columns": cache["unmapped"],
    }


@dataclass
class LoadReport:
    registered: int = 0
    reused: int = 0
    skipped_series: int = 0
    written: int = 0
    unchanged: int = 0
    rejected: list[str] = field(default_factory=list)
    # 列映射缓存命中的列数
    column_cache_hits: int = 0
    # 编码与缓存都没认出来的列，每项带 sheet / 指纹 / 列名 / 编码，等人确认一次
    unmapped_columns: list[dict] = field(default_factory=list)

    def line(self) -> str:
        out = (f"指标 新登记 {self.registered} / 复用 {self.reused} / 跳过停用 {self.skipped_series}；"
               f"观测 写入 {self.written:,} / 值未变跳过 {self.unchanged:,} / 拒绝 {len(self.rejected)}")
        if self.column_cache_hits or self.unmapped_columns:
            out += (f"；列映射 缓存命中 {self.column_cache_hits} / "
                    f"待确认 {len(self.unmapped_columns)}")
        return out


def load(session, data_or_path, *, entered_by: str, variety: str = "SN",
         dry_run: bool = False, progress=None, register_new: bool = True,
         catalog_only: bool = False) -> LoadReport:
    """整本入库。所有观测仍逐条走 `record()` 闸门，只是共用一份预取的最新值索引。

    `register_new=False` 时拒绝库里没见过的编码，只更新已登记序列。**无人值守的自动导入
    必须用这个模式**：有人把编码敲错一位、或凭空编一个，自动登记会让一条查无实据的序列
    混进事实层，而没有人在确认页看过它。有人盯着的手工导入才用默认的 True。

    `catalog_only=True` 只登记指标、不写任何观测。用于冷启动建目录：在终端里按目录整批
    选中、**只导一天**，文件很小但带着全部编码、名称、单位、频率。导完平台就知道有哪些
    序列了，之后的数据工作簿直接从平台取编码，不必在终端界面里一个个找。
    """
    from tin.ingest.record import RecordError, latest_index, record
    from tin.models import Indicator

    report = LoadReport()
    everything = read_workbook(data_or_path)
    wanted = [s for s in everything if not discontinued(s)]
    report.skipped_series = len(everything) - len(wanted)

    ids, _ = resolve_ids(session, wanted)
    # 编码认不出的列，交给列映射学习缓存兜一次底。试算不计命中数。
    cache = column_cache(session, wanted, ids, count_hit=not dry_run)
    report.column_cache_hits = len(cache["mapped"])
    report.unmapped_columns = cache["unmapped"]
    registered: dict[str, object] = {}
    known = set()
    for series in wanted:
        spec = derive_spec(series, variety, series_id=ids[series.code])
        existing = session.get(Indicator, spec.series_id)
        if existing is None and not register_new:
            report.rejected.append(
                f"{series.code}: 编码未登记，自动导入模式不新建指标（名称「{series.name[:30]}」）"
                "——确认一次列映射后，同格式的导出以后会自动命中")
            continue
        if existing is None:
            report.registered += 1
            if not dry_run:
                existing = Indicator(
                    series_id=spec.series_id, name=spec.name, variety=spec.variety,
                    category=spec.category, caliber=spec.caliber.dump(), unit=spec.unit,
                    source=spec.source, source_url=spec.source_url, frequency=spec.frequency,
                    fetch_mode=spec.fetch_mode, phase="P1", vendor_code=series.code,
                    default_entry_time=entry_time_for(series),
                    # 没有人为这条指标的口径背书过：它是照着文件里的名称自动推出来的。
                    # 研究员复核后把 owner 改成自己，就算认领了。
                    owner=UNVERIFIED,
                )
                session.add(existing)
        else:
            report.reused += 1
            if not dry_run and not existing.vendor_code:
                existing.vendor_code = series.code
        registered[series.code] = existing
        known.add(series.code)
    if not dry_run:
        session.flush()
    if dry_run:
        report.written = sum(len(s.points) for s in wanted)
        return report
    if catalog_only:
        session.commit()
        return report

    wanted = [s for s in wanted if s.code in known]
    cache = latest_index(session, [ids[s.code] for s in wanted])
    for n, series in enumerate(wanted, 1):
        for obs in observations(series, entered_by, registered.get(series.code), variety,
                                series_id=ids[series.code]):
            try:
                row = record(session, obs, cache=cache)
            except RecordError as exc:
                report.rejected.append(f"{ids[series.code]}: {exc}")
                break  # 同一序列后续点必然同样被拒，不必刷屏
            report.written += row is not None
            report.unchanged += row is None
        session.flush()
        # 最后一条必定回调：否则不足 20 条的工作簿进度永远停在 0%
        if progress and (n % 20 == 0 or n == len(wanted)):
            progress(n, len(wanted), report)
    session.commit()
    return report
