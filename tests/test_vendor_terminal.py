"""数据商终端导出工作簿的整体导入（docs/数据进出方案.md §3）。"""

import io
from datetime import date, datetime, time

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from tin.caliber.dictionary import ContractKind, PriceType, SpotSource, StockScope, WeightBasis
from tin.config import SHANGHAI
from tin.ingest import vendor_terminal as vt
from tin.ingest.record import latest_index, record
from tin.models import Indicator, Observation
from tin.schemas.caliber import Caliber
from tin.schemas.observation import ObservationIn


def make_workbook(sheets: dict[str, list[list]]) -> io.BytesIO:
    """按 SMM 终端的四行表头布局造一本工作簿。"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


SAMPLE = {
    "Sheet1": [
        ["指标名称", "SHFE: 锡: 主力合约: 收盘价: 日度", "SHFE: 锡: 连二合约: 结算价: 日度",
         "（停）SHFE: 锡: 仓单日报分仓库_旧库: 期货: 日度"],
        ["指标Id", "a1001", "a1002", "a1003"],
        ["单位", "元/吨", "元/吨", "吨"],
        ["频率", "日", "日", "日"],
        [datetime(2026, 9, 18), 407490, 405000, 12],
        [datetime(2026, 9, 21), 409000, 408000, 13],
    ],
    "Sheet2": [
        ["指标名称", "SMM: 1#锡-平均价: 日度", "SMM: 锡矿矿端平衡: 缅甸矿进口量: 月度",
         "SMM: 锡锭社会库存: 总库存: 周度", "SMM: 锡现货买卖强度指数: 日度"],
        ["指标Id", "s20002960", "a10005168", "a10020744", "a10021649"],
        ["单位", "元/吨", "实物吨", "吨", "-"],
        ["频率", "日", "月", "周", "日"],
        [datetime(2026, 9, 18), 408200, 15000, 9100, 51.2],
        [datetime(2026, 9, 21), 409350, None, 8800, "-"],
    ],
    "说明": [["本表为说明页"], ["无数据"]],
}


@pytest.fixture
def series():
    return {s.code: s for s in vt.read_workbook(make_workbook(SAMPLE))}


def test_reads_the_four_row_header_layout(series):
    """指标Id 是对齐主键：改了显示名也不会错位，这是这条通道不做名称猜测的底气。"""
    assert set(series) == {"a1001", "a1002", "a1003", "s20002960", "a10005168", "a10020744", "a10021649"}
    spot = series["s20002960"]
    assert spot.unit == "元/吨" and spot.frequency == "日"
    assert spot.points == {date(2026, 9, 18): 408200.0, date(2026, 9, 21): 409350.0}


def test_sheets_without_the_header_are_skipped(series):
    """「说明」这类目录页没有四行表头，整张跳过而不是硬解析出一堆空列。"""
    assert all(s.sheet != "说明" for s in series.values())


def test_blank_and_dash_cells_are_not_zero(series):
    """SMM 用空单元格和「-」表示无值，绝不能当成 0 写进事实层。"""
    assert series["a10005168"].points == {date(2026, 9, 18): 15000.0}
    assert series["a10021649"].points == {date(2026, 9, 18): 51.2}


def test_duplicate_vendor_ids_are_taken_once():
    """同一 Id 在表里既有短名又有全名，只能留一条，否则同一序列会被登记两次。"""
    sheets = {"S": [
        ["指标名称", "SHFE: 锡锭库存总计: 周度", "SHFE: 锡: 库存周报: 本周库存小计: 周度"],
        ["指标Id", "a10024418", "a10024418"],
        ["单位", "吨", "吨"],
        ["频率", "周", "周"],
        [datetime(2026, 9, 18), 7059, 7059],
    ]}
    got = vt.read_workbook(make_workbook(sheets))
    assert [s.code for s in got] == ["a10024418"]


def test_discontinued_columns_are_flagged(series):
    assert vt.discontinued(series["a1003"])
    assert not vt.discontinued(series["a1001"])


def test_override_binds_to_the_already_registered_series(series):
    """已登记的现货价不另建一份，否则 BASIS 会指着一条空序列算不出数。"""
    assert vt.SMM.series_id("s20002960") == "SMM.SN.spot.1"
    assert vt.SMM.series_id("a1001") == "SMM.a1001"


def test_caliber_is_derived_from_the_name_hierarchy(series):
    cal = vt.derive_caliber(series["a1001"]).dump()
    assert cal["price_type"] == PriceType.收盘价
    assert cal["contract_kind"] == ContractKind.主力合约
    assert cal["time_type"] == "交易时点"

    assert vt.derive_caliber(series["a1002"]).dump()["contract_kind"] == ContractKind.连二连续
    assert vt.derive_caliber(series["a10020744"]).dump()["stock_scope"] == StockScope.社会库存
    assert vt.derive_caliber(series["a10005168"]).dump()["weight_basis"] == WeightBasis.实物吨
    assert vt.derive_caliber(series["s20002960"]).dump()["spot_source"] == SpotSource.SMM
    # SMM 自采数据是发布时点，不是交易时点
    assert vt.derive_caliber(series["a10020744"]).dump()["time_type"] == "发布时点"


def test_unknown_spot_market_stays_blank_rather_than_inventing_a_value():
    """个旧不在口径字典里。塞一个字典外的值会让同时点守卫失效，宁可留空并标注。"""
    sheets = {"S": [
        ["指标名称", "SMM: 个旧1#锡 - 平均价: 日度", "SMM: 长江 1#锡-平均价: 日度"],
        ["指标Id", "s22795417", "s20123078"],
        ["单位", "元/吨", "元/吨"],
        ["频率", "日", "日"],
        [datetime(2026, 9, 21), 408000, 409000],
    ]}
    by_code = {s.code: s for s in vt.read_workbook(make_workbook(sheets))}
    gejiu = vt.derive_caliber(by_code["s22795417"]).dump()
    assert "spot_source" not in gejiu or gejiu["spot_source"] is None
    assert "个旧" in gejiu["note"]
    assert vt.derive_caliber(by_code["s20123078"]).dump()["spot_source"] == SpotSource.长江有色


def test_load_registers_indicators_and_writes_observations(session):
    report = vt.load(session, make_workbook(SAMPLE), entered_by="测试")
    assert report.skipped_series == 1  # 停用列
    assert report.rejected == []
    assert report.written == 10  # 两表各列的有效点之和，停用列不计

    ind = session.get(Indicator, "SMM.a1001")
    assert ind is not None and ind.vendor_code == "a1001" and ind.fetch_mode == "manual"
    assert ind.source == vt.SMM.label

    rows = session.scalars(select(Observation).where(Observation.series_id == "SMM.a1001")).all()
    assert {r.value for r in rows} == {407490.0, 409000.0}
    # 日期单元格没有时刻，按发布方的常规发布时刻补齐，并在 note 里说明补齐规则
    assert {r.as_of.astimezone(SHANGHAI).strftime("%H:%M") for r in rows} == {"15:00"}
    assert "补齐" in rows[0].note


def test_load_is_idempotent(session):
    first = vt.load(session, make_workbook(SAMPLE), entered_by="测试")
    second = vt.load(session, make_workbook(SAMPLE), entered_by="测试")
    assert first.written > 0
    assert second.written == 0
    assert second.unchanged == first.written
    assert second.registered == 0


def test_registered_caliber_wins_over_name_derivation(session):
    """已登记指标以登记口径入库，不用按名称猜的那个。

    否则钢联「锡锭：库存：中国」会被推成交易所库存，而它登记的是社会库存——
    一条本该入库的数据会被闸门挡住，挡的还是推导规则自己的错。
    """
    session.add(Indicator(
        series_id="SMM.a1001", name="占位", variety="SN", category="价格",
        caliber={"time_type": "交易时点", "price_type": "结算价"},
        unit="元/吨", source="SMM 终端", frequency="日", fetch_mode="manual", phase="P1"))
    session.flush()

    report = vt.load(session, make_workbook(SAMPLE), entered_by="测试")
    assert report.rejected == []
    rows = session.scalars(select(Observation).where(Observation.series_id == "SMM.a1001")).all()
    assert len(rows) == 2
    assert rows[0].caliber_snapshot["price_type"] == "结算价", "入库快照记的是登记口径"


def test_load_still_goes_through_the_record_gate(session):
    """闸门没被绕过：停用的指标一律拒收，批量通道不是后门。"""
    session.add(Indicator(
        series_id="SMM.a1001", name="占位", variety="SN", category="价格",
        caliber={"time_type": "交易时点", "price_type": "收盘价"},
        unit="元/吨", source="SMM 终端", frequency="日", fetch_mode="manual",
        phase="P1", status="停用"))
    session.flush()

    report = vt.load(session, make_workbook(SAMPLE), entered_by="测试")
    assert any("已停用" in r for r in report.rejected)
    assert session.scalars(select(Observation).where(Observation.series_id == "SMM.a1001")).all() == []


def test_batch_cache_behaves_exactly_like_the_uncached_gate(session):
    """`record()` 的缓存参数只为省查询，判重与修订号必须与逐条查询完全一致。"""
    session.add(Indicator(
        series_id="T.X", name="测试", variety="SN", category="价格", caliber={"time_type": "发布时点"},
        unit="元/吨", source="测试", frequency="日", fetch_mode="auto", phase="P1"))
    session.flush()
    at = datetime(2026, 9, 21, 15, tzinfo=SHANGHAI)

    def obs(value):
        return ObservationIn(series_id="T.X", value=value, as_of=at,
                             caliber=Caliber(time_type="发布时点"), source="测试")

    assert record(session, obs(1.0)).revision == 0
    assert record(session, obs(1.0)) is None          # 值未变
    assert record(session, obs(2.0)).revision == 1    # 值变了，追加修订

    cache = latest_index(session, ["T.X"])
    assert record(session, obs(2.0), cache=cache) is None
    bumped = record(session, obs(3.0), cache=cache)
    assert bumped.revision == 2
    # 缓存要跟着更新，否则同一批里的后续点会拿到过期的修订号
    assert record(session, obs(3.0), cache=cache) is None
    session.flush()
    assert len(session.scalars(select(Observation).where(Observation.series_id == "T.X")).all()) == 3


def test_a_workbook_without_the_smm_layout_is_rejected_loudly():
    """格式不符要明确报错，不能静默返回空——静默返回空正是网页导入那个「四个 0」的病根。"""
    sheets = {"S": [["日期", "随便一列"], [datetime(2026, 9, 21), 1]]}
    with pytest.raises(vt.VendorFormatError):
        vt.read_workbook(make_workbook(sheets))


# ---------- 网页整体导入通道 ----------

def test_detects_terminal_layout_without_reading_everything():
    """探测必须发生在单元格上限之前，所以只看前四行的首列标签。"""
    assert vt.detect(make_workbook(SAMPLE).getvalue()) is not None
    plain = {"S": [["日期", "SMM 1# 锡现货均价 [SMM.SN.spot.1]"], [datetime(2026, 9, 21), 409000]]}
    assert vt.detect(make_workbook(plain).getvalue()) is None
    assert vt.detect(b"not an xlsx at all") is None


def test_summary_answers_what_a_human_can_actually_judge(session):
    """几十万行没人逐行看得完，确认页给的是批次级信息。"""
    s = vt.summarize(session, make_workbook(SAMPLE))
    assert s["mode"] == "vendor_terminal"
    assert s["new_indicators"] == 5
    assert [r["series_id"] for r in s["reused_indicators"]] == ["SMM.SN.spot.1"]
    assert s["skipped_discontinued"] == 1
    assert s["points"] == 10
    assert (s["first_date"], s["last_date"]) == ("2026-09-18", "2026-09-21")
    assert s["caliber_notes"] == []
    assert sum(s["categories"].values()) == s["new_indicators"]


def test_summary_surfaces_where_registration_and_derivation_disagree(session):
    """登记口径与按名称推导的不一致时，摆出来给人看——但以登记为准，不阻断。

    人按供应商编码做的绑定比按名称的推导权威；两者打架只说明推导规则有问题，
    或者那条登记绑错了编码。两种都该让人知道，而不是静默选一个。
    """
    session.add(Indicator(
        series_id="SMM.a1001", name="占位", variety="SN", category="价格",
        caliber={"time_type": "交易时点", "price_type": "结算价"},  # 文件里那列是「收盘价」
        unit="元/吨", source="SMM 终端", frequency="日", fetch_mode="manual", phase="P1"))
    session.flush()
    s = vt.summarize(session, make_workbook(SAMPLE))
    assert any("price_type 登记为 结算价，按名称推导为 收盘价" in note for note in s["caliber_notes"])


def test_staging_token_cannot_escape_the_staging_directory(tmp_path, monkeypatch):
    """令牌直接参与拼文件名，必须挡住路径穿越。"""
    from tin.config import settings
    from tin.ingest import import_jobs

    monkeypatch.setattr(settings, "staging_dir", tmp_path)
    for bad in ("../../etc/passwd", "a" * 31, "../" + "a" * 29, "ABCDEF" + "0" * 26):
        with pytest.raises(ValueError, match="令牌"):
            import_jobs.staging_path(bad)

    token = import_jobs.stage(b"x")
    assert import_jobs.staging_path(token).parent == tmp_path


def test_job_status_never_leaks_the_staging_token(tmp_path, monkeypatch):
    from tin.config import settings
    from tin.ingest import import_jobs

    monkeypatch.setattr(settings, "staging_dir", tmp_path)
    token = import_jobs.stage(make_workbook(SAMPLE).getvalue())
    job_id = import_jobs.create_terminal_job(token, "测试员", "tin.xlsx")
    state = import_jobs.snapshot(job_id)
    assert "token" not in state
    assert state["state"] == "running" and state["percent"] == 0


def test_background_job_imports_and_reports(tmp_path, monkeypatch, session):
    """后台任务跑完要把结果留在状态里，并写一条审计。"""
    from sqlalchemy.orm import sessionmaker

    from tin.config import settings
    from tin.ingest import import_jobs
    from tin.models import AuditLog

    monkeypatch.setattr(settings, "staging_dir", tmp_path)
    monkeypatch.setattr("tin.db.SessionLocal", sessionmaker(session.get_bind(), expire_on_commit=False))

    token = import_jobs.stage(make_workbook(SAMPLE).getvalue())
    job_id = import_jobs.create_terminal_job(token, "测试员", "tin.xlsx")
    import_jobs.execute(job_id)

    state = import_jobs.snapshot(job_id)
    assert state["state"] == "done", state.get("error")
    assert state["registered"] == 5 and state["written"] == 10
    assert state["percent"] == 100
    # 暂存文件跑完即删，不留副本
    assert not import_jobs.staging_path(token).exists()

    log = session.scalars(select(AuditLog).where(AuditLog.action == "导入")).all()
    assert len(log) == 1 and log[0].detail["written"] == 10


def test_background_job_records_failure_instead_of_swallowing_it(tmp_path, monkeypatch):
    from tin.config import settings
    from tin.ingest import import_jobs

    monkeypatch.setattr(settings, "staging_dir", tmp_path)
    token = import_jobs.stage("这不是一个 xlsx".encode())
    job_id = import_jobs.create_terminal_job(token, "测试员", "坏文件.xlsx")
    import_jobs.execute(job_id)
    state = import_jobs.snapshot(job_id)
    assert state["state"] == "failed" and state["error"]


# ---------- 钢联 Mysteel 版式 ----------

MYSTEEL_SAMPLE = {
    "Sheet1": [  # 少一行「更新时间」——钢联自己的表头行数就不固定
        ["钢联数据"],
        ["指标名称", "锡精矿：40%Sn：加工费：云南（日）", "锡精矿：60%Sn：加工费：广西（日）"],
        ["单位", "元/吨", "元/吨"],
        ["指标编码", "ID01538256", "ID01538259"],
        ["频度", "日", "日"],
        ["指标描述", "·", "·"],
        [datetime(2026, 12, 31), "#N/A", "#N/A"],
        [datetime(2026, 9, 22), 17500, 12000],
        [datetime(2026, 9, 19), 17300, "#N/A"],
    ],
    "Sheet2": [
        ["钢联数据"],
        ["指标名称", "锡锭：库存：中国（周）", "LME：锡：期货库存（日）", "SHFE：锡：主力合约：结算价（日）"],
        ["单位", "吨", "吨", "元/吨"],
        ["指标编码", "ID01517441", "FU00015899", "FU00016040"],
        ["频度", "周", "日", "日"],
        ["指标描述", "样本覆盖全国主要产区", "·", "·"],
        ["更新时间", "2026-09-22 11:35:08", "·", "·"],
        [datetime(2026, 9, 18), 8968, 4780, 408500],
    ],
}


@pytest.fixture
def mysteel():
    return {s.code: s for s in vt.read_workbook(make_workbook(MYSTEEL_SAMPLE))}


def test_mysteel_layout_is_detected_and_parsed(mysteel):
    """钢联的标签、顺序、表头行数都和 SMM 不同，靠「首列是日期」定位数据起点。"""
    assert vt.detect(make_workbook(MYSTEEL_SAMPLE).getvalue()) is vt.MYSTEEL
    assert set(mysteel) == {"ID01538256", "ID01538259", "ID01517441", "FU00015899", "FU00016040"}
    tc = mysteel["ID01538256"]
    assert tc.unit == "元/吨" and tc.frequency == "日"
    assert tc.points == {date(2026, 9, 22): 17500.0, date(2026, 9, 19): 17300.0}


def test_mysteel_na_marker_is_not_zero(mysteel):
    """钢联用 #N/A 表示无值，当成 0 会直接污染事实层。"""
    assert date(2026, 12, 31) not in mysteel["ID01538256"].points
    assert mysteel["ID01538259"].points == {date(2026, 9, 22): 12000.0}


def test_mysteel_description_lands_in_the_caliber_note(mysteel):
    note = vt.derive_caliber(mysteel["ID01517441"]).dump()["note"]
    assert "钢联 ID01517441" in note and "样本覆盖全国主要产区" in note


def test_tc_grade_is_read_from_the_name(mysteel):
    """40%Sn 与 60%Sn 的加工费数值不可比，品位是必须落下来的口径维度。"""
    assert vt.derive_caliber(mysteel["ID01538256"]).dump()["tc_grade"] == "40度"
    assert vt.derive_caliber(mysteel["ID01538259"]).dump()["tc_grade"] == "60度"
    # 加工费不是价格类型，不能硬塞一个
    assert "price_type" not in vt.derive_caliber(mysteel["ID01538256"]).dump()


def test_exchange_stock_and_survey_stock_are_not_the_same_scope(mysteel):
    """LME 的期货库存是交易所库存，钢联自采的锡锭库存是社会库存，两者不可混算。"""
    assert vt.derive_caliber(mysteel["FU00015899"]).dump()["stock_scope"] == "交易所库存"
    assert vt.derive_caliber(mysteel["ID01517441"]).dump()["stock_scope"] == "社会库存"
    assert vt.derive_caliber(mysteel["FU00015899"]).dump()["time_type"] == "交易时点"
    assert vt.derive_caliber(mysteel["ID01517441"]).dump()["time_type"] == "发布时点"


def test_mysteel_overrides_bind_the_three_registered_indicators(mysteel, session):
    """这三个编码在库里已有登记，必须归到原 series_id，否则判断读的还是空序列。"""
    assert vt.MYSTEEL.series_id("ID01538256") == "MYSTEEL.SN.TC.YN40"
    assert vt.MYSTEEL.series_id("ID01517441") == "MYSTEEL.SN.stock.social"
    assert vt.MYSTEEL.series_id("FU00082529") == "ICDX.SN.volume"
    assert vt.MYSTEEL.series_id("ID01538259") == "MYSTEEL.ID01538259"


def test_registered_entry_time_is_used_for_known_indicators(mysteel, session):
    """已登记指标的发布时刻以登记为准，否则同一序列走导入和走录入页会落在不同时刻。"""
    from tin.models import Indicator

    tc = session.get(Indicator, "MYSTEEL.SN.TC.YN40")
    assert tc.default_entry_time == time(11, 0)
    obs = vt.observations(mysteel["ID01538256"], "测试", registered=tc)
    assert {o.as_of.astimezone(SHANGHAI).strftime("%H:%M") for o in obs} == {"11:00"}
    # 未登记的按发布方的常规时刻
    assert vt.entry_time_for(mysteel["ID01538259"]) == time(15, 0)


def test_summary_counts_points_dated_in_the_future(session):
    """年频/季频按期末标注，会出现晚于今天的点——不拦，但要让人知道有多少。"""
    s = vt.summarize(session, make_workbook(MYSTEEL_SAMPLE))
    assert s["vendor"] == "钢联终端"
    assert s["future_points"] == 0  # 12-31 那行全是 #N/A，没有真值
    assert any("ICDX" not in m for m in s["frequency_mismatches"]) or s["frequency_mismatches"]


def test_mysteel_load_writes_through_the_same_gate(session):
    report = vt.load(session, make_workbook(MYSTEEL_SAMPLE), entered_by="测试")
    assert report.rejected == []
    assert report.reused == 2  # TC 与锡锭库存已登记；ICDX 不在这个样本里
    rows = session.scalars(
        select(Observation).where(Observation.series_id == "MYSTEEL.SN.TC.YN40")).all()
    assert {r.value for r in rows} == {17500.0, 17300.0}
    assert rows[0].caliber_snapshot["tc_grade"] == "40度"


# ---------- 编码归属 ----------

def test_unknown_vendor_code_registers_a_new_indicator(session):
    """库里没见过的编码 → 自动登记新指标并入库，不需要改代码。"""
    sheets = {"S": [
        ["指标名称", "SMM: 锡: 某个全新指标: 日度"],
        ["指标Id", "a99999001"],
        ["单位", "元/吨"],
        ["频率", "日"],
        [datetime(2026, 9, 21), 12345],
    ]}
    report = vt.load(session, make_workbook(sheets), entered_by="测试")
    assert (report.registered, report.written, report.rejected) == (1, 1, [])
    ind = session.get(Indicator, "SMM.a99999001")
    assert ind.vendor_code == "a99999001" and ind.fetch_mode == "manual"


def test_db_binding_beats_the_hardcoded_override_table(session):
    """研究员在库里把编码登记到可读 series_id 上之后，导入必须归到那里。

    只查代码里的 OVERRIDES 的话，会另建一个 MYSTEEL.<编码>，同一条序列无声裂成两个。
    """
    from datetime import time as _t

    session.add(Indicator(
        series_id="MYSTEEL.SN.newthing", name="某新指标", variety="SN", category="需求",
        caliber={"time_type": "发布时点"}, unit="%", source="Mysteel", frequency="月",
        fetch_mode="manual", phase="P1", vendor_code="ID09999001", default_entry_time=_t(15, 0)))
    session.flush()

    sheets = {"S": [
        ["钢联数据"], ["指标名称", "锡：某新指标：中国（月）"], ["单位", "%"],
        ["指标编码", "ID09999001"], ["频度", "月"], ["指标描述", "·"],
        [datetime(2026, 8, 31), 62.5],
    ]}
    report = vt.load(session, make_workbook(sheets), entered_by="测试")
    assert report.registered == 0 and report.reused == 1
    assert session.get(Indicator, "MYSTEEL.ID09999001") is None, "不该另建一份"
    rows = session.scalars(
        select(Observation).where(Observation.series_id == "MYSTEEL.SN.newthing")).all()
    assert [r.value for r in rows] == [62.5]


def test_one_code_bound_to_two_indicators_is_reported_not_guessed(session):
    """同一编码被登记到两个指标上，不替人选，报出来让人先合并。"""
    from datetime import time as _t

    for sid in ("MYSTEEL.SN.dupA", "MYSTEEL.SN.dupB"):
        session.add(Indicator(
            series_id=sid, name=sid, variety="SN", category="需求",
            caliber={"time_type": "发布时点"}, unit="%", source="Mysteel", frequency="月",
            fetch_mode="manual", phase="P1", vendor_code="ID09999002", default_entry_time=_t(15, 0)))
    session.flush()

    sheets = {"S": [
        ["钢联数据"], ["指标名称", "锡：重复编码：中国（月）"], ["单位", "%"],
        ["指标编码", "ID09999002"], ["频度", "月"], ["指标描述", "·"],
        [datetime(2026, 8, 31), 1.0],
    ]}
    summary = vt.summarize(session, make_workbook(sheets))
    assert any("同时登记在" in c for c in summary["code_conflicts"])


def test_plain_table_will_not_register_anything_new(session):
    """逐行通道不会自动登记：未登记的代码或名称一律归入异常，不静默造指标。"""
    from tin.ingest.excel_importer import build_preview

    tagged = make_workbook({"S": [["日期", "某新指标 [SMM.a88888]"], ["2026-09-21", 999]]})
    pv = build_preview(session, tagged.getvalue(), "t.xlsx", actor="测试", variety="SN")
    assert pv["ready_count"] == 0
    assert "未在指标登记表中找到" in pv["columns"][0]["error"]

    named = make_workbook({"S": [["日期", "某个从没登记过的指标"], ["2026-09-21", 999]]})
    pv = build_preview(session, named.getvalue(), "t.xlsx", actor="测试", variety="SN")
    assert pv["ready_count"] == 0
    assert pv["columns"][0]["error"] == "系统未登记该指标"
    assert session.get(Indicator, "SMM.a88888") is None


# ---------- 数据可信性 ----------

def test_uncalculated_plugin_formulas_fail_loudly(tmp_path):
    """数据终端的 Excel 插件写入的是公式，结果要 Excel 保存过才缓存进文件。

    openpyxl 不算公式，读到 None——整本文件会被当成"全是缺失值"静默导入 0 条。
    这种失败必须吵：安静地导入 0 条，比报错危险得多。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["指标名称", "某插件序列"])
    ws.append(["指标Id", "a12345"])
    ws.append(["单位", "元/吨"])
    ws.append(["频率", "日"])
    ws.append([datetime(2026, 9, 22), '=MSD("a12345",TODAY())'])
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(vt.VendorFormatError, match="公式尚未计算"):
        vt.read_workbook(buf.getvalue())


def test_unattended_mode_refuses_unknown_codes(session):
    """无人值守的自动导入不能自动登记：敲错一位或凭空编一个编码，
    会让一条查无实据的序列混进事实层，而没有人在确认页看过它。"""
    sheets = {"S": [
        ["指标名称", "SMM: 锡: 查无此项: 日度"],
        ["指标Id", "a00000000"],
        ["单位", "元/吨"],
        ["频率", "日"],
        [datetime(2026, 9, 21), 1],
    ]}
    report = vt.load(session, make_workbook(sheets), entered_by="定时任务", register_new=False)
    assert report.registered == 0 and report.written == 0
    assert any("自动导入模式不新建指标" in r for r in report.rejected)
    assert session.get(Indicator, "SMM.a00000000") is None


def test_auto_registered_indicators_are_marked_unverified(session):
    """自动登记的指标没有人为它的口径背书，要标出来等人认领。"""
    sheets = {"S": [
        ["指标名称", "SMM: 锡: 全新序列: 日度"], ["指标Id", "a00000001"],
        ["单位", "元/吨"], ["频率", "日"], [datetime(2026, 9, 21), 1],
    ]}
    vt.load(session, make_workbook(sheets), entered_by="测试")
    assert session.get(Indicator, "SMM.a00000001").owner == vt.UNVERIFIED
    # 已登记的指标不会被这个标记覆盖
    assert session.get(Indicator, "SMM.SN.spot.1").owner != vt.UNVERIFIED


def test_catalog_only_registers_without_writing_observations(session):
    """冷启动建目录：终端按目录整批只导一天，文件很小但带着全部编码与元数据。

    编码无法从我们这边反推，只能由终端吐出来——所以要有一条"只要目录不要数据"的路。
    """
    report = vt.load(session, make_workbook(SAMPLE), entered_by="测试", catalog_only=True)
    assert report.registered == 5
    assert report.written == 0
    ind = session.get(Indicator, "SMM.a1001")
    assert ind is not None and ind.vendor_code == "a1001"
    assert session.scalars(select(Observation).where(Observation.series_id == "SMM.a1001")).all() == []

    # 目录建好之后，无人值守的数据导入就不再需要新建指标
    data = vt.load(session, make_workbook(SAMPLE), entered_by="定时任务", register_new=False)
    assert data.registered == 0 and data.written == 10 and data.rejected == []
