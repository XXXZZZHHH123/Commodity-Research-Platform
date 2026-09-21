"""Excel 自定义导出：不出数保护、低频对齐、研判过期标记、合规留痕与模板治理。"""

import io
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from tests.conftest import D, enter_spot
from tin.compute.engine import compute_day
from tin.config import SHANGHAI
from tin.export.excel import COMPLIANCE_TEXT, Column, build_export, field_catalog
from tin.export.templates import TemplateError, delete_template, list_templates, save_template
from tin.jobs.seed import seed_judgment
from tin.models import ExportTemplate, Indicator, TradingDay


def sheet(data: bytes, index: int = 0):
    wb = load_workbook(io.BytesIO(data))
    return wb[wb.sheetnames[index]]


def cells(ws, col: int) -> list:
    return [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]


COLS = [Column("trade_date", "交易日", "meta"),
        Column("SHFE.SN.main.settle", "主力结算价", "observation"),
        Column("BASIS", "基差", "derived")]


# ---------- 字段目录 ----------

def test_catalog_is_built_from_the_registry_not_hardcoded(session):
    groups = {g["name"]: g["fields"] for g in field_catalog(session, "SN")}
    assert "派生指标" in groups and "研判" in groups
    all_fields = [f["field"] for fields in groups.values() for f in fields]
    assert "SMM.SN.spot.1" in all_fields
    assert "BASIS" in all_fields
    assert not any(f.startswith("SHFE.SN.26") for f in all_fields), "逐合约序列按月换约，不进长时序宽表"


def test_catalog_skips_disabled_indicators(session):
    session.get(Indicator, "SMM.SN.spot.1").status = "停用"
    session.commit()
    fields = [f["field"] for g in field_catalog(session, "SN") for f in g["fields"]]
    assert "SMM.SN.spot.1" not in fields


# ---------- 不出数保护 ----------

def test_blocked_cell_is_empty_with_reason_comment(loaded):
    """基差因时点错配被阻断时，Excel 里必须留空并写明原因，不能是 0、也不能无声空着。"""
    enter_spot(loaded, 401000, datetime(2026, 9, 17, 11, 30, tzinfo=SHANGHAI))
    compute_day(loaded, D)
    ws = sheet(build_export(loaded, "SN", COLS, actor="张三")[0])

    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "2026-09-18")
    cell = ws.cell(row=row, column=3)
    assert cell.value is None
    assert "阻断不出数" in cell.comment.text and "不在同一交易日" in cell.comment.text


def test_uncomputed_cell_is_distinguished_from_blocked(loaded):
    """从未跑过计算 ≠ 算过但被阻断，两者批注必须不同。"""
    ws = sheet(build_export(loaded, "SN", COLS, actor="张三")[0])
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "2026-09-18")
    assert ws.cell(row=row, column=3).comment.text.startswith("[未计算]")


def test_missing_input_is_not_labelled_as_blocked(loaded):
    """缺输入是没数可算，阻断是算得出但口径对不上——两种空格含义不同，批注必须分开。"""
    compute_day(loaded, D)  # 现货未录入，基差为 missing_input
    ws = sheet(build_export(loaded, "SN", COLS, actor="张三")[0])
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "2026-09-18")
    text = ws.cell(row=row, column=3).comment.text
    assert text.startswith("[缺少输入]") and "阻断" not in text


def test_ok_derived_value_is_written_with_number_format(loaded):
    enter_spot(loaded, 406000, datetime(2026, 9, 18, 11, 30, tzinfo=SHANGHAI))
    compute_day(loaded, D)
    ws = sheet(build_export(loaded, "SN", COLS, actor="张三")[0])
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "2026-09-18")
    cell = ws.cell(row=row, column=3)
    assert cell.value == 640 and cell.comment is None
    assert cell.number_format == "#,##0"


# ---------- 低频对齐，禁止前向填充 ----------

def _weekly_scenario(session):
    """9/10–9/18 为交易日，周库存只在 9/11 与 9/18 两个周五有值。"""
    from tin.ingest.record import record
    from tin.schemas.caliber import Caliber
    from tin.schemas.observation import ObservationIn
    for day, seq in ((10, 166), (11, 167), (14, 168), (15, 169), (16, 170)):
        if session.get(TradingDay, f"2026-09-{day}") is None:
            session.add(TradingDay(trade_date=f"2026-09-{day}", year=2026, year_seq=seq))
    for day, value in ((11, 7652), (18, 7059)):
        record(session, ObservationIn(
            series_id="SHFE.SN.stock.weekly", value=value,
            as_of=datetime(2026, 9, day, 15, tzinfo=SHANGHAI),
            caliber=Caliber(time_type="交易时点", stock_scope="交易所库存"), source="SHFE"))
    session.commit()
    return [Column("trade_date", "交易日", "meta"),
            Column("SHFE.SN.stock.weekly", "周库存", "observation")]


def test_weekly_series_is_blank_on_non_publish_days(loaded):
    """覆盖区间之内的非发布日必须留空并注明，不得沿用前值。"""
    ws = sheet(build_export(loaded, "SN", _weekly_scenario(loaded), actor="张三")[0])
    assert [v for v in cells(ws, 2) if v is not None] == [7059, 7652]

    inside = [ws.cell(row=r, column=2) for r in range(2, ws.max_row + 1)
              if ws.cell(row=r, column=2).value is None
              and "2026-09-11" < ws.cell(row=r, column=1).value < "2026-09-18"]
    assert inside and all("非发布日" in c.comment.text for c in inside)


def test_dates_before_a_series_starts_are_not_annotated(loaded):
    """长跨度导出时，指标尚未开始的那段留白不逐格批注，改在说明页写明覆盖区间。"""
    data = build_export(loaded, "SN", _weekly_scenario(loaded), actor="张三")[0]
    ws = sheet(data)
    before = next(ws.cell(row=r, column=2) for r in range(2, ws.max_row + 1)
                  if ws.cell(row=r, column=1).value == "2026-09-10")
    assert before.value is None and before.comment is None

    text = "\n".join(str(c.value) for row in sheet(data, 1).iter_rows() for c in row if c.value)
    assert "本区间数据覆盖" in text and "2026-09-11 至 2026-09-18" in text


def test_monthly_value_on_non_trading_day_is_carried_to_next_trade_date(loaded):
    """海关月度数据 as_of 常落在非交易日的月末，必须顺延到下一交易日，并注明原始时点。"""
    from tin.ingest.record import record
    from tin.schemas.caliber import Caliber
    from tin.schemas.observation import ObservationIn
    loaded.add(TradingDay(trade_date="2026-09-21", year=2026, year_seq=175))
    record(loaded, ObservationIn(
        series_id="CUSTOMS.2609.import.MM", value=12000,
        as_of=datetime(2026, 9, 20, 23, 59, tzinfo=SHANGHAI),  # 周日
        caliber=Caliber(time_type="交易时点", weight_basis="实物吨",
                        import_code="精矿(HS2609)", country="缅甸"),
        source="人工", entered_by="张三", note="海关月报"))
    loaded.commit()

    cols = [Column("trade_date", "交易日", "meta"),
            Column("CUSTOMS.2609.import.MM", "缅甸进口", "observation")]
    ws = sheet(build_export(loaded, "SN", cols, actor="张三")[0])
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "2026-09-21")
    cell = ws.cell(row=row, column=2)
    assert cell.value == 12000
    assert "顺延对齐" in cell.comment.text and "2026-09-20" in cell.comment.text


# ---------- 研判列 ----------

def test_judgment_column_always_carries_version_and_expiry(loaded):
    """孤立的基调词一旦脱离系统就会被当成当下结论，必须绑定版本与过期天数。"""
    seed_judgment(loaded)
    loaded.commit()
    cols = [Column("trade_date", "交易日", "meta"), Column("judgment_tone", "基调", "judgment")]
    ws = sheet(build_export(loaded, "SN", cols, actor="张三")[0])
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "2026-09-18")
    text = ws.cell(row=row, column=2).value
    assert text.startswith("区间震荡（v1，") and "已过期 3 天" in text


def test_judgment_column_is_empty_before_the_first_version(loaded):
    seed_judgment(loaded)
    judgment = loaded.scalars(select(__import__("tin.models", fromlist=["Judgment"]).Judgment)).one()
    judgment.written_at = D + timedelta(days=5)
    loaded.commit()
    cols = [Column("trade_date", "交易日", "meta"), Column("judgment_tone", "基调", "judgment")]
    ws = sheet(build_export(loaded, "SN", cols, actor="张三")[0])
    assert all(v is None for v in cells(ws, 2))


# ---------- 表头、排版与合规 ----------

def test_header_carries_codes_for_lossless_round_trip(loaded):
    from tin.ingest.excel_importer import match_header

    ws = sheet(build_export(loaded, "SN", COLS, actor="张三")[0])
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert header[1] == "沪锡主力合约结算价 [SHFE.SN.main.settle]"
    assert ws.freeze_panes == "B2"

    spec = match_header(header[1], list(loaded.scalars(select(Indicator))))
    assert spec.match_kind == "code" and not spec.needs_choice


def test_compliance_notice_lives_outside_the_data_grid(loaded):
    """合规提示走页眉与说明页，数据表第 1 行必须是纯表头，保证机器可读。"""
    data = build_export(loaded, "SN", COLS, actor="张三")[0]
    ws = sheet(data)
    assert ws.cell(row=1, column=1).value == "交易日"
    assert COMPLIANCE_TEXT in ws.oddHeader.center.text

    dictionary = sheet(data, 1)
    assert dictionary.title == "口径与合规说明"
    assert COMPLIANCE_TEXT in dictionary.cell(row=1, column=1).value


def test_dictionary_sheet_lists_caliber_and_source_of_every_column(loaded):
    dictionary = sheet(build_export(loaded, "SN", COLS, actor="张三")[0], 1)
    text = "\n".join(str(c.value) for row in dictionary.iter_rows() for c in row if c.value)
    assert "SHFE.SN.main.settle" in text and "结算价" in text
    assert "BASIS" in text and "容差" in text
    assert "阻断不出数" in text and "不沿用前值" in text


def test_export_reports_commercial_sources_for_audit(loaded):
    cols = COLS + [Column("SMM.SN.spot.1", "现货", "observation")]
    _, meta = build_export(loaded, "SN", cols, actor="张三")
    assert meta["commercial_sources"] == ["SMM"]


def test_disabled_indicator_is_skipped_not_fatal(loaded):
    loaded.get(Indicator, "SMM.SN.spot.1").status = "停用"
    loaded.commit()
    cols = COLS + [Column("SMM.SN.spot.1", "现货", "observation")]
    data, meta = build_export(loaded, "SN", cols, actor="张三")
    assert meta["skipped"] == ["SMM.SN.spot.1"]
    header = [c.value for c in sheet(data)[1]]
    assert not any("SMM.SN.spot.1" in str(h) for h in header)
    assert "已停用" in "\n".join(str(c.value) for row in sheet(data, 1).iter_rows()
                              for c in row if c.value)


def test_sort_order_and_range(loaded):
    ws_desc = sheet(build_export(loaded, "SN", COLS, actor="张三", sort_order="desc")[0])
    ws_asc = sheet(build_export(loaded, "SN", COLS, actor="张三", sort_order="asc")[0])
    assert cells(ws_desc, 1) == sorted(cells(ws_desc, 1), reverse=True)
    assert cells(ws_asc, 1) == sorted(cells(ws_asc, 1))

    custom = build_export(loaded, "SN", COLS, actor="张三", date_range="custom",
                          start_date="2026-09-18", end_date="2026-09-18")[1]
    assert custom["rows"] == 1


# ---------- 模板治理 ----------

def payload(**over) -> dict:
    return {"name": "晨报核心", "columns": [{"field": "trade_date", "label": "交易日", "kind": "meta"}],
            **over}


def test_first_template_becomes_default(session):
    saved = save_template(session, "SN", payload())
    assert saved["is_default"] is True


def test_column_kind_is_required(session):
    with pytest.raises(TemplateError, match="kind"):
        save_template(session, "SN", payload(columns=[{"field": "BASIS", "label": "基差"}]))


def test_duplicate_name_is_rejected(session):
    save_template(session, "SN", payload())
    with pytest.raises(TemplateError, match="同名模板"):
        save_template(session, "SN", payload())


def test_only_one_default_per_variety(session):
    first = save_template(session, "SN", payload())
    save_template(session, "SN", payload(name="周度平衡", is_default=True))
    rows = list_templates(session, "SN")
    assert [r["is_default"] for r in rows].count(True) == 1
    assert next(r for r in rows if r["id"] == first["id"])["is_default"] is False


def test_deleting_default_promotes_the_most_recent_survivor(session):
    """默认模板必须可删，删除后由剩余模板顶上，避免出现删不掉的死角。"""
    first = save_template(session, "SN", payload())
    second = save_template(session, "SN", payload(name="周度平衡"))
    assert first["is_default"] and not second["is_default"]

    delete_template(session, "SN", first["id"])
    rows = list_templates(session, "SN")
    assert len(rows) == 1 and rows[0]["id"] == second["id"] and rows[0]["is_default"] is True


def test_deleting_last_template_leaves_empty_state(session):
    saved = save_template(session, "SN", payload())
    result = delete_template(session, "SN", saved["id"])
    assert result["templates"] == []
    assert session.scalars(select(ExportTemplate)).first() is None
